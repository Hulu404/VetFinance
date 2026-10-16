"""Excel file of a report: numbers are numbers (roubles with kopecks, percents in percent units)."""

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.exports.document import (
    BRAND,
    Cell,
    Days,
    Document,
    FigureRef,
    Pct,
    Pp,
    Rub,
    Table,
    cell_of,
    sources_text,
)

MONEY_FORMAT = "#,##0.00"
PCT_FORMAT = "0.0"
DATE_FORMAT = "DD.MM.YYYY"
ACCENT = "8A0F3A"
HEADER_FILL = PatternFill("solid", fgColor="F8E1E9")
KOPECKS = Decimal(100)
WIDTH_PER_UNIT = 6


def _write(
    ws: Worksheet, row: int, column: int, value: Cell, *, bold: bool = False, indent: int = 0
) -> None:
    cell = ws.cell(row=row, column=column)
    if isinstance(value, Rub):
        cell.value = (Decimal(value.kopecks) / KOPECKS).quantize(Decimal("0.01"))
        cell.number_format = MONEY_FORMAT
    elif isinstance(value, Pct | Pp):
        cell.value = value.value
        cell.number_format = PCT_FORMAT
    elif isinstance(value, Days):
        cell.value = value.value
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = DATE_FORMAT
    else:
        cell.value = value
    if bold:
        cell.font = Font(bold=True)
    if indent:
        cell.alignment = Alignment(indent=indent)


def _header_block(ws: Worksheet, doc: Document) -> int:
    """Brand, clinic, title, period, times, headline, reasons; returns the next free row."""
    ws.cell(row=1, column=1, value=BRAND).font = Font(bold=True, size=14, color=ACCENT)
    ws.cell(row=2, column=1, value=doc.clinic)
    ws.cell(row=3, column=1, value=f"{doc.title} · {doc.period_line}").font = Font(bold=True)
    ws.cell(row=4, column=1, value=f"Сформировано: {doc.made_at}")
    ws.cell(row=5, column=1, value=f"Данные на: {doc.data_as_of or 'нет данных'}")
    row = 6
    if doc.headline:
        ws.cell(row=row, column=1, value=doc.headline)
        row += 1
    for reason in doc.reasons:
        ws.cell(row=row, column=1, value=f"Внимание: {reason}")
        row += 1
    return row + 1


def _table(ws: Worksheet, table: Table, start: int) -> int:
    for col, title in enumerate(table.columns, start=1):
        cell = ws.cell(row=start, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    row = start + 1
    for r in table.rows:
        for col, value in enumerate(r.cells, start=1):
            _write(ws, row, col, value, bold=r.bold, indent=r.indent if col == 1 else 0)
        row += 1
    for col, width in enumerate(table.widths, start=1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(
            ws.column_dimensions[letter].width or 0, width * WIDTH_PER_UNIT
        )
    return row


def _composition(wb: Workbook, figures: tuple[FigureRef, ...]) -> None:
    ws = wb.create_sheet("Состав")
    for col, title in enumerate(("Строка", "Слагаемое", "Значение"), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    row = 2
    for ref in figures:
        fig = ref.figure
        _write(ws, row, 1, ref.label, bold=True)
        _write(ws, row, 3, cell_of(fig.value, fig.unit), bold=True)
        row += 1
        for item in fig.breakdown:
            _write(ws, row, 2, item.label)
            _write(ws, row, 3, cell_of(item.value, item.unit))
            row += 1
        for item in fig.details or []:
            _write(ws, row, 2, f"{item.label} (справочно)")
            _write(ws, row, 3, cell_of(item.value, item.unit))
            row += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 18


def _sources(wb: Workbook, doc: Document) -> None:
    ws = wb.create_sheet("Источники")
    for col, title in enumerate(("Показатель", "Источники данных", "Пояснение"), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    row = 2
    for reason in doc.reasons:
        _write(ws, row, 1, "Отчёт неполный")
        _write(ws, row, 3, reason)
        row += 1
    for ref in doc.figures:
        _write(ws, row, 1, ref.label)
        _write(ws, row, 2, sources_text(ref.figure))
        _write(ws, row, 3, ref.figure.note)
        row += 1
    for note in doc.footnotes:
        _write(ws, row, 1, "Примечание")
        _write(ws, row, 3, note)
        row += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 100


def render_xlsx(doc: Document) -> bytes:
    wb = Workbook()
    first = wb.active
    assert first is not None  # noqa: S101 (a new workbook always has an active sheet)
    for index, table in enumerate(doc.tables):
        ws = first if index == 0 else wb.create_sheet(table.title[:31])
        ws.title = table.title[:31]
        start = _header_block(ws, doc) if index == 0 else 1
        _table(ws, table, start)
    _composition(wb, doc.figures)
    _sources(wb, doc)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
