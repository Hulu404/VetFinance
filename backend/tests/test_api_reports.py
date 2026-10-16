"""Stage 3 report endpoints over the in-memory repository loaded with golden set 4 (tests/golden/set04, read only).

Numbers are compared with expected.json (computed independently of this code); the contract itself (shapes,
identities inside the response, Figure rules, validation, authentication, export, rate limit) is checked here.
"""

import io
import json
from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.schemas import FigureOut
from app.secrets import FernetBox
from tests.conftest import FakeVetmanagerApi, build_app, make_token
from tests.set04_loader import EXPECTED, load_set04

OPENAPI = Path(__file__).resolve().parents[1] / "openapi.json"
HEADERS = {"Authorization": f"Bearer {make_token()}"}
SOURCE_ORDER = ["bank", "vetmanager", "bot", "manual", "calc"]
ROW_KEYS = ["revenue", "variable", "marginal_profit", "fixed", "taxes", "obligations", "net_profit"]
ROW_KINDS = ["revenue", "expense", "subtotal", "expense", "expense", "expense", "total"]
RATING_KEYS = ("status", "tone", "deviation_pp", "deviation_rub")
MONTHS = ["2026-07", "2026-08", "2026-09"]
AS_OF_STOCK = EXPECTED["stock"]["as_of"]
PATHS = [
    "/v1/reports/pnl",
    "/v1/reports/norm",
    "/v1/reports/stock",
    "/v1/reports/stock/idle",
    "/v1/reports/pnl/export?format=xlsx",
]


@pytest.fixture(scope="module")
def client_e() -> Iterator[TestClient]:
    app = build_app(load_set04(), FakeVetmanagerApi(), FernetBox(Fernet.generate_key().decode()))
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def _fresh_export_budget(client_e: TestClient) -> None:
    """The module client is shared: every test starts with a full budget of exports."""
    client_e.app.state.export_limiter.reset()  # type: ignore[attr-defined]


def get(client: TestClient, path: str, **params: Any) -> dict[str, Any]:
    r = client.get(path, params=params, headers=HEADERS)
    assert r.status_code == 200, r.text
    body: dict[str, Any] = r.json()
    return body


def figures(node: Any, path: str = "") -> Iterator[tuple[str, dict[str, Any]]]:
    if isinstance(node, dict):
        if {"value", "unit", "breakdown", "source", "note"} <= node.keys():
            yield path, node
        for k, v in node.items():
            yield from figures(v, f"{path}.{k}")
    elif isinstance(node, list):
        for i, v in enumerate(node):
            yield from figures(v, f"{path}[{i}]")


def rows(d: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {r["key"]: r for r in d["rows"]}


def value(row: dict[str, Any]) -> int:
    v = row["amount"]["value"]
    assert isinstance(v, int)
    return v


def labelled(figure: dict[str, Any]) -> dict[str, int | float]:
    return {b["label"]: b["value"] for b in figure["breakdown"]}


# ------------------------------------------------------------------ authentication
@pytest.mark.parametrize("path", PATHS)
def test_reports_need_a_token(client: TestClient, path: str) -> None:
    r = client.get(path)
    assert r.status_code == 401
    assert r.json()["error"]["code"] == "unauthorized"


# ------------------------------------------------------------------ Figure rules in every report
@pytest.mark.parametrize(
    ("path", "params"),
    [
        ("/v1/reports/pnl", {"view": "clinic"}),
        ("/v1/reports/pnl", {"view": "direction"}),
        ("/v1/reports/pnl", {"month": "2026-08", "view": "direction"}),
        ("/v1/reports/norm", {}),
        ("/v1/reports/norm", {"month": "2026-08"}),
        ("/v1/reports/norm", {"month": "2026-09", "as_of": "2026-09-15"}),
        ("/v1/reports/stock", {}),
        ("/v1/reports/stock/idle", {}),
    ],
)
def test_every_figure_has_source_note_and_adds_up(
    client_e: TestClient, path: str, params: dict[str, str]
) -> None:
    found = list(figures(get(client_e, path, **params)))
    assert len(found) >= 5
    for where, fig in found:
        FigureOut.model_validate(fig)
        assert fig["note"].strip(), where
        assert fig["source"], where
        assert fig["source"] == sorted(set(fig["source"]), key=SOURCE_ORDER.index), where
        assert fig["breakdown"] or (fig["unit"] == "rub" and fig["value"] == 0), where
        units = {b["unit"] for b in fig["breakdown"]}
        if fig["unit"] in ("rub", "count") and units == {fig["unit"]}:
            assert sum(b["value"] for b in fig["breakdown"]) == fig["value"], where
        if fig["unit"] in ("rub", "count", "days"):
            assert isinstance(fig["value"], int), where


# ------------------------------------------------------------------ P&L
@pytest.mark.parametrize("month", MONTHS)
def test_pnl_equals_the_golden_set(client_e: TestClient, month: str) -> None:
    d = get(client_e, "/v1/reports/pnl", month=month)
    exp = EXPECTED["pnl"][month]
    assert [r["key"] for r in d["rows"]] == ROW_KEYS
    assert [r["kind"] for r in d["rows"]] == ROW_KINDS
    for e in exp["rows"]:
        got = rows(d)[e["key"]]
        assert value(got) == e["amount"]["value"]
        assert list(labelled(got["amount"]).items()) == list(e["amount"]["breakdown"].items())
        assert got["amount"]["source"] == e["amount"]["source"]
        assert got["share_pct"]["value"] == e["share_pct"]["value"]
    assert d["unclassified"]["value"] == exp["unclassified"]["amount"]
    assert d["unclassified_count"] == exp["unclassified"]["count"]
    assert (d["status"], d["partial_reasons"]) == (exp["status"], exp["partial_reasons"])
    assert d["period"] == {
        "month": month,
        "status": "closed",
        "from": f"{month}-01",
        "to": None,
        "first_month": "2026-07",
    }
    assert (d["month"], d["view"], d["as_of"]) == (month, "clinic", "2026-09-30")
    assert d["data_as_of"] is not None


def test_pnl_without_month_is_the_last_closed_month(client_e: TestClient) -> None:
    assert get(client_e, "/v1/reports/pnl")["month"] == "2026-09"
    # on the 15th of September the last closed month is August
    assert get(client_e, "/v1/reports/pnl", as_of="2026-09-15")["month"] == "2026-08"


def test_pnl_identity_and_structure(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/pnl", month="2026-09")
    by = rows(d)
    assert by["obligations"]["label"] == "Обязательства и амортизация"
    assert "warning" not in by["obligations"]
    expenses = sum(value(by[k]) for k in ("variable", "fixed", "taxes", "obligations"))
    assert value(by["revenue"]) - expenses == value(by["net_profit"])
    assert value(by["marginal_profit"]) == value(by["revenue"]) - value(by["variable"])
    kids = by["fixed"]["children"]
    assert [c["key"] for c in kids] == ["fixed_payroll", "fixed_other"]
    assert sum(value(c) for c in kids) == value(by["fixed"])
    assert {c["kind"] for c in kids} == {"child"}
    assert all("children" not in r for k, r in by.items() if k != "fixed")
    assert "directions" not in d and "common_variable" not in d
    assert "128 544,93" in d["headline"]


def test_pnl_share_is_a_percent_figure_with_two_rouble_rows(client_e: TestClient) -> None:
    for r in get(client_e, "/v1/reports/pnl")["rows"]:
        share = r["share_pct"]
        assert share["unit"] == "pct"
        assert [b["label"] for b in share["breakdown"]] == ["Сумма строки", "Выручка"]


@pytest.mark.parametrize("month", MONTHS)
def test_pnl_direction_view_matches_the_golden_set_and_bridges(
    client_e: TestClient, month: str
) -> None:
    d = get(client_e, "/v1/reports/pnl", month=month, view="direction")
    assert d["view"] == "direction"
    assert d["rows"] == get(client_e, "/v1/reports/pnl", month=month)["rows"]
    exp = EXPECTED["direction"][month]
    got = {x["direction"]: x for x in d["directions"]}
    for e in exp["directions"]:
        g = got[e["direction"]]
        assert g["revenue"]["value"] == e["revenue"]["value"]
        assert g["direct_costs"]["value"] == e["direct_costs"]
        assert g["margin"]["value"] == e["margin"]
        assert g["margin_pct"]["value"] == e["margin_pct"]
        assert g["revenue_share_pct"]["value"] == e["revenue_share_pct"]
    assert d["common_variable"]["value"] == exp["common_variable"]["value"]
    by, (services, goods) = rows(d), d["directions"]
    common = d["common_variable"]["value"]
    assert services["revenue"]["value"] + goods["revenue"]["value"] == value(by["revenue"])
    direct = services["direct_costs"]["value"] + goods["direct_costs"]["value"]
    assert direct + common == value(by["variable"])
    margins = services["margin"]["value"] + goods["margin"]["value"]
    assert margins - common == value(by["marginal_profit"])


def test_pnl_running_month(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/pnl", month="2026-09", as_of="2026-09-15")
    exp = {r["key"]: r for r in EXPECTED["pnl"]["2026-09@2026-09-15"]["rows"]}
    assert d["period"]["status"] == "open" and d["period"]["to"] == "2026-09-15"
    assert d["as_of"] == "2026-09-15"
    assert {k: value(r) for k, r in rows(d).items()} == {
        k: e["amount"]["value"] for k, e in exp.items()
    }
    assert d["headline"].startswith("К 15 сентября клиника заработала")
    assert "налоги и взносы, уже уплаченные" in rows(d)["taxes"]["amount"]["note"].lower()
    assert "за весь месяц" in rows(d)["obligations"]["amount"]["note"]


def test_pnl_month_without_data_is_empty_and_has_no_shares(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/pnl", month="2026-06")
    assert d["status"] == "empty" and "headline" not in d
    assert all(value(r) == 0 for r in d["rows"])
    assert all("share_pct" not in r for r in d["rows"])
    assert d["period"]["first_month"] == "2026-07"


def test_pnl_for_a_clinic_without_reports_data_is_empty(client: TestClient) -> None:
    d = get(client, "/v1/reports/pnl", month="2026-09")
    assert d["status"] == "empty" and d["unclassified_count"] == 0
    assert rows(d)["obligations"]["warning"] == "no_obligations"


# ------------------------------------------------------------------ norm
@pytest.mark.parametrize("month", MONTHS)
def test_norm_items_equal_the_golden_set(client_e: TestClient, month: str) -> None:
    d = get(client_e, "/v1/reports/norm", month=month)
    exp = EXPECTED["norm"][month]["items"]
    assert [i["key"] for i in d["items"]] == [e["key"] for e in exp]
    for got, e in zip(d["items"], exp, strict=True):
        assert (got["status"], got["tone"], got["memo"], got["kind"]) == (
            e["status"],
            e["tone"],
            e["memo"],
            e["kind"],
        )
        assert got["actual_pct"]["value"] == e["actual_pct"]
        assert list(labelled(got["actual_pct"]).items()) == list(e["amount_breakdown"].items())
        if e["status"] == "no_norm":
            assert "norm" not in got and "deviation_pp" not in got and "deviation_rub" not in got
            continue
        assert got["deviation_rub"]["value"] == e["deviation_rub"]
        assert got["deviation_pp"]["value"] == e["deviation_pp"]
        assert got["deviation_pp"]["unit"] == "pp" and got["deviation_pp"]["delta"] is True
        assert got["deviation_rub"]["delta"] is True
        assert got["norm"]["source"] == e["norm"]["source"]
        assert got["norm"]["min_pct"] == e["norm"]["min_pct"]
        assert got["norm"]["max_pct"] == e["norm"]["max_pct"]


def test_norm_shares_cover_the_whole_revenue_and_deviations_add_up(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/norm", month="2026-09")
    revenue = value(rows(get(client_e, "/v1/reports/pnl"))["revenue"])
    total = sum(
        b["value"] for i in d["items"] if not i["memo"] for b in i["actual_pct"]["breakdown"]
    )
    assert total == revenue
    for i in d["items"]:
        if i.get("status") in ("above", "below"):
            parts = sum(b["value"] for b in i["deviation_rub"]["breakdown"])
            assert parts == i["deviation_rub"]["value"]
        if i.get("status") == "ok":
            assert i["deviation_rub"]["breakdown"] == [] and i["deviation_rub"]["details"]
    memo = [i["key"] for i in d["items"] if i["memo"]]
    assert memo == ["writeoff"]


def test_norm_open_month_has_no_ratings_but_keeps_the_borders(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/norm", month="2026-09", as_of="2026-09-15")
    assert d["period"]["status"] == "open"
    assert d["headline"] == EXPECTED["norm"]["2026-09@2026-09-15"]["headline_open"]
    for i in d["items"]:
        assert not any(k in i for k in RATING_KEYS), i["key"]
    assert d["items"][0]["norm"]["max_pct"] == 28.0
    short = get(client_e, "/v1/reports/norm", month="2026-09", as_of="2026-09-05")
    assert short["partial_reasons"] == ["short_period"] and short["status"] == "partial"


def test_norm_headline_of_a_closed_month_mentions_the_deviations(client_e: TestClient) -> None:
    text = get(client_e, "/v1/reports/norm", month="2026-09")["headline"]
    assert "1 594,79" in text and "21 114,75" in text and "в расходы месяца оно не входит" in text


def test_norm_of_a_month_without_revenue_is_empty(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/norm", month="2026-06")
    assert d["status"] == "empty" and d["items"] == [] and "headline" not in d


def test_norm_without_benchmarks_says_so(client_e: TestClient) -> None:
    repo = load_set04()
    repo.benchmarks = []
    app = build_app(repo, FakeVetmanagerApi(), FernetBox(Fernet.generate_key().decode()))
    with TestClient(app) as c:
        d = get(c, "/v1/reports/norm", month="2026-09")
    assert d["partial_reasons"] == ["no_benchmark"] and d["status"] == "partial"
    assert all(i["status"] == "no_norm" for i in d["items"])
    assert d["headline"] == "Нормы для вашей клиники пока нет, но доли статей ниже."


# ------------------------------------------------------------------ stock
def test_stock_equals_the_golden_set(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/stock", as_of=AS_OF_STOCK)
    exp = EXPECTED["stock"]
    assert d["month"] == "2026-09" and d["status"] == "ok"
    assert d["frozen"]["value"] == exp["frozen"]["value"]
    assert labelled(d["frozen"]) == exp["frozen"]["breakdown"]
    got_rows = [
        (r["item"], r["expires_on"], r["days_left"], r["qty"], r["amount"]["value"])
        for r in d["expiring"]
    ]
    want_rows = [
        (r["item"], r["expires_on"], r["days_left"], r["qty"], r["amount"])
        for r in exp["expiring_60d"]["rows"]
    ]
    assert got_rows == want_rows
    assert d["expiring_60d"]["value"] == exp["expiring_60d"]["value"]
    assert all(r["days_left"] < 0 for r in d["expired_items"])
    assert d["expired_on_hand"]["value"] == exp["expired_on_hand"]["value"]
    assert d["written_off_month"]["value"] == exp["written_off_month"]["value"]
    written = [
        (w["item"], w["date"], w["qty"], w["reason"], w["amount"]["value"])
        for w in d["written_off_items"]
    ]
    assert written == [
        (w["item"], w["date"], w["qty"], w["reason"], w["amount"]) for w in exp["written_off_items"]
    ]
    assert d["written_off_norm"]["value"] == exp["written_off_norm"]["value"] == 878_525
    assert d["turnover_days"]["value"] == exp["turnover_days"]["value"]
    assert d["turnover_days"]["unit"] == "days" and d["norm_turnover_days"] == 45
    assert d["stock_value"]["value"] == exp["stock_value"]["value"]
    assert d["frozen_positions"]["value"] == exp["frozen_positions"]["value"]
    assert [
        (b["key"], b["amount"]["value"], b["positions"]["value"]) for b in d["idle_buckets"]
    ] == [(b["key"], b["amount"], b["positions"]) for b in exp["idle_buckets"]]
    assert "expiring_at_risk" not in d and "missing_expiry_count" not in d  # reserved fields


def test_stock_write_off_norm_equals_the_norm_row(client_e: TestClient) -> None:
    s = get(client_e, "/v1/reports/stock", as_of=AS_OF_STOCK)
    n = get(client_e, "/v1/reports/norm", month="2026-09")
    row = next(i for i in n["items"] if i["key"] == "writeoff")
    assert s["written_off_norm"]["value"] == row["norm"]["breakdown"][-1]["value"]
    expired = sum(w["amount"]["value"] for w in s["written_off_items"] if w["reason"] == "expired")
    assert expired == sum(b["value"] for b in row["actual_pct"]["breakdown"])


def test_stock_month_defaults_to_the_last_ended_month(client_e: TestClient) -> None:
    assert get(client_e, "/v1/reports/stock", as_of="2026-09-15")["month"] == "2026-08"
    assert get(client_e, "/v1/reports/stock", as_of="2026-09-30")["month"] == "2026-09"
    august = get(client_e, "/v1/reports/stock", as_of="2026-09-30", month="2026-08")
    assert august["month"] == "2026-08" and august["stock_value"]["value"] == 55_354_600


def test_stock_of_a_clinic_without_stock_data_is_empty(client: TestClient) -> None:
    d = get(client, "/v1/reports/stock")
    assert d["status"] == "empty" and "headline" not in d
    assert d["stock_value"]["value"] == 0 and d["expiring"] == [] and d["written_off_items"] == []
    assert d["turnover_days"]["value"] == 0 and d["norm_turnover_days"] == 45
    idle = get(client, "/v1/reports/stock/idle")
    assert idle["items"] == [] and idle["total"]["value"] == 0 and idle["next_cursor"] is None


def test_stock_idle_pages_equal_the_golden_pages(client_e: TestClient) -> None:
    for case in EXPECTED["stock"]["idle_pages"]:
        cursor: str | None = None
        for page in case["pages"]:
            params: dict[str, Any] = {
                "bucket": case["bucket"],
                "limit": case["limit"],
                "as_of": AS_OF_STOCK,
            }
            if cursor:
                params["cursor"] = cursor
            d = get(client_e, "/v1/reports/stock/idle", **params)
            assert [
                (
                    i["item"],
                    i["last_movement_on"],
                    i["idle_days"],
                    i["bucket"],
                    i["qty"],
                    i["amount"]["value"],
                )
                for i in d["items"]
            ] == [
                (
                    i["item"],
                    i["last_movement_on"],
                    i["idle_days"],
                    i["bucket"],
                    i["qty"],
                    i["amount"],
                )
                for i in page["items"]
            ]
            assert d["total"]["value"] == case["total"] and d["total"]["unit"] == "count"
            assert (d["next_cursor"] is not None) == page["has_next"]
            cursor = d["next_cursor"]


# ------------------------------------------------------------------ validation
@pytest.mark.parametrize("bad", ["garbage", "2026-13", "2026-9", "26-09", "2026-09-01", "0000-09", ""])
def test_invalid_month_is_422(client: TestClient, bad: str) -> None:
    for path in ("/v1/reports/pnl", "/v1/reports/norm", "/v1/reports/stock"):
        r = client.get(path, params={"month": bad}, headers=HEADERS)
        assert r.status_code == 422
        assert r.json()["error"]["details"]["fields"]["month"] == "invalid"


def test_future_month_is_422(client: TestClient) -> None:
    for path in ("/v1/reports/pnl", "/v1/reports/norm", "/v1/reports/stock"):
        r = client.get(path, params={"month": "2026-10"}, headers=HEADERS)
        assert r.status_code == 422
        assert r.json()["error"]["details"]["fields"] == {"month": "in_future"}
    ok = client.get("/v1/reports/pnl", params={"month": "2026-09"}, headers=HEADERS)
    assert ok.status_code == 200
    moved = client.get(
        "/v1/reports/pnl", params={"month": "2026-10", "as_of": "2026-10-02"}, headers=HEADERS
    )
    assert moved.status_code == 200 and moved.json()["period"]["status"] == "open"


@pytest.mark.parametrize("bad", ["", "all", "CLINIC"])
def test_invalid_view_is_422(client: TestClient, bad: str) -> None:
    r = client.get("/v1/reports/pnl", params={"view": bad}, headers=HEADERS)
    assert r.status_code == 422 and "view" in r.json()["error"]["details"]["fields"]


def test_invalid_as_of_is_422(client: TestClient) -> None:
    for path in ("/v1/reports/pnl", "/v1/reports/norm", "/v1/reports/stock"):
        r = client.get(path, params={"as_of": "nope"}, headers=HEADERS)
        assert r.status_code == 422 and r.json()["error"]["code"] == "validation"


@pytest.mark.parametrize(
    "params", [{"bucket": "x"}, {"cursor": "abc"}, {"limit": 0}, {"limit": 51}, {"limit": "x"}]
)
def test_invalid_idle_params_are_422(client: TestClient, params: dict[str, Any]) -> None:
    r = client.get("/v1/reports/stock/idle", params=params, headers=HEADERS)
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation"


def test_clinic_id_parameter_is_ignored(client_e: TestClient) -> None:
    base = get(client_e, "/v1/reports/pnl")
    other = get(client_e, "/v1/reports/pnl", clinic_id="00000000-0000-0000-0000-000000000000")
    assert base == other


def test_the_same_request_gives_identical_numbers(client_e: TestClient) -> None:
    for path in ("/v1/reports/pnl", "/v1/reports/norm", "/v1/reports/stock"):
        assert get(client_e, path) == get(client_e, path)


# ------------------------------------------------------------------ export
def export(client: TestClient, name: str, fmt: str, **params: Any) -> Any:
    r = client.get(f"/v1/reports/{name}/export", params={"format": fmt, **params}, headers=HEADERS)
    assert r.status_code == 200, r.text[:300]
    return r


def sheet_rows(content: bytes) -> list[list[Any]]:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    return [list(row) for row in wb.worksheets[0].iter_rows(values_only=True)]


def find(rows_: list[list[Any]], label: str) -> list[Any]:
    return next(r for r in rows_ if label in r)


def kopecks(cell: Any) -> int:
    return int(round(Decimal(str(cell)) * 100))


def pdf_text(content: bytes) -> str:
    text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(content)).pages)
    return "".join(ch for ch in text if not ch.isspace() and ch != " ")


def rub_digits(value_: int) -> str:
    rub, kop = divmod(abs(value_), 100)
    return f"{rub}" + (f",{kop:02d}" if kop else "")


@pytest.mark.parametrize(
    ("name", "params", "filename"),
    [
        ("pnl", {"month": "2026-09"}, "vetfinance-pnl-2026-09.xlsx"),
        ("norm", {"month": "2026-08"}, "vetfinance-norm-2026-08.xlsx"),
        ("stock", {"as_of": AS_OF_STOCK}, "vetfinance-stock-2026-09.xlsx"),
    ],
)
def test_export_headers(
    client_e: TestClient, name: str, params: dict[str, str], filename: str
) -> None:
    x = export(client_e, name, "xlsx", **params)
    assert x.headers["content-type"].split(";")[0].endswith("spreadsheetml.sheet")
    assert x.headers["content-disposition"] == f'attachment; filename="{filename}"'
    assert x.content[:2] == b"PK"
    p = export(client_e, name, "pdf", **params)
    assert p.headers["content-type"].split(";")[0] == "application/pdf"
    assert p.headers["content-disposition"] == f'attachment; filename="{filename[:-4]}pdf"'
    assert p.content[:5] == b"%PDF-"


def test_pnl_xlsx_holds_the_numbers_of_the_json_as_numbers(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/pnl", month="2026-09", view="direction")
    content = export(client_e, "pnl", "xlsx", month="2026-09").content
    first = sheet_rows(content)
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Вся клиника", "Услуги и товары", "Состав", "Источники"]
    for r in d["rows"] + [c for row in d["rows"] for c in row.get("children") or []]:
        row = find(first, r["label"])
        numbers = [c for c in row if isinstance(c, int | float | Decimal)]
        assert kopecks(numbers[0]) == r["amount"]["value"], r["label"]
        assert abs(float(numbers[1]) - r["share_pct"]["value"]) < 1e-9, r["label"]
    directions = {x["label"]: x for x in d["directions"]}
    second = [list(row) for row in wb["Услуги и товары"].iter_rows(values_only=True)]
    for label, x in directions.items():
        row = find(second, label)
        assert kopecks(row[1]) == x["revenue"]["value"] and kopecks(row[4]) == x["margin"]["value"]
    sources = [list(row) for row in wb["Источники"].iter_rows(values_only=True)]
    assert any("Ветменеджер" in str(r[1]) for r in sources)
    composition = [list(row) for row in wb["Состав"].iter_rows(values_only=True)]
    assert any(r[1] == "Услуги" and kopecks(r[2]) == 64_132_070 for r in composition)


def test_norm_xlsx_holds_shares_deviations_and_statuses(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/norm", month="2026-09")
    first = sheet_rows(export(client_e, "norm", "xlsx", month="2026-09").content)
    for i in d["items"]:
        row = find(first, i["label"])
        assert abs(row[2] - i["actual_pct"]["value"]) < 1e-9
        if "deviation_rub" in i:
            assert kopecks(row[7]) == i["deviation_rub"]["value"]
            assert abs(row[6] - i["deviation_pp"]["value"]) < 1e-9
    assert find(first, "Аренда и коммуналка")[5] == "выше нормы"
    assert find(first, "Остальные статьи")[5] == "нормы нет"
    assert (
        find(first, "Выручка месяца")[1] == Decimal("1098156.10")
        or kopecks(find(first, "Выручка месяца")[1]) == 109_815_610
    )


def test_stock_xlsx_and_pdf_hold_the_totals(client_e: TestClient) -> None:
    d = get(client_e, "/v1/reports/stock", as_of=AS_OF_STOCK)
    content = export(client_e, "stock", "xlsx", as_of=AS_OF_STOCK).content
    wb = load_workbook(io.BytesIO(content), data_only=True)
    cells = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row]
    money = {kopecks(c) for c in cells if isinstance(c, int | float | Decimal)}
    for key in ("frozen", "expiring_60d", "written_off_month", "stock_value", "written_off_norm"):
        assert d[key]["value"] in money, key
    assert d["turnover_days"]["value"] in {c for c in cells if isinstance(c, int)}
    text = pdf_text(export(client_e, "stock", "pdf", as_of=AS_OF_STOCK).content)
    for key in ("frozen", "expiring_60d", "written_off_month", "stock_value"):
        assert rub_digits(d[key]["value"]) in text, key
    for r in d["expiring"]:
        assert r["item"].replace(" ", "") in text
        assert rub_digits(r["amount"]["value"]) in text


@pytest.mark.parametrize(
    ("name", "params"), [("pnl", {"month": "2026-09"}), ("norm", {"month": "2026-09"})]
)
def test_pdf_text_holds_labels_and_amounts_in_cyrillic(
    client_e: TestClient, name: str, params: dict[str, str]
) -> None:
    d = get(client_e, f"/v1/reports/{name}", **params)
    text = pdf_text(export(client_e, name, "pdf", **params).content)
    assert "ВетФинанс" in text and "Берёзка" in text
    if name == "pnl":
        for r in d["rows"] + [c for row in d["rows"] for c in row.get("children") or []]:
            assert r["label"].replace(" ", "") in text, r["label"]
            assert rub_digits(r["amount"]["value"]) in text, r["label"]
        assert "Услуги" in text and "Товары" in text
    else:
        for i in d["items"]:
            assert i["label"].replace(" ", "") in text, i["label"]
            if "deviation_rub" in i and i["deviation_rub"]["value"]:
                assert rub_digits(i["deviation_rub"]["value"]) in text


def test_export_of_an_empty_month_still_works(client_e: TestClient) -> None:
    for name in ("pnl", "norm"):
        assert export(client_e, name, "xlsx", month="2026-06").content[:2] == b"PK"
        assert export(client_e, name, "pdf", month="2026-06").content[:5] == b"%PDF-"


def test_export_of_a_clinic_without_stock_data_works(client: TestClient) -> None:
    assert export(client, "stock", "xlsx").content[:2] == b"PK"
    assert export(client, "stock", "pdf").content[:5] == b"%PDF-"


@pytest.mark.parametrize("bad", ["", "csv", "XLSX", "doc"])
def test_export_unknown_format_is_422(client: TestClient, bad: str) -> None:
    r = client.get("/v1/reports/pnl/export", params={"format": bad}, headers=HEADERS)
    assert r.status_code == 422 and "format" in r.json()["error"]["details"]["fields"]
    missing = client.get("/v1/reports/pnl/export", headers=HEADERS)
    assert missing.json()["error"]["details"]["fields"]["format"] == "required"


def test_export_unknown_report_is_404(client: TestClient) -> None:
    r = client.get("/v1/reports/nothing/export", params={"format": "pdf"}, headers=HEADERS)
    assert r.status_code == 404 and r.json()["error"]["code"] == "not_found"


def test_export_validates_the_month(client: TestClient) -> None:
    r = client.get(
        "/v1/reports/pnl/export", params={"format": "xlsx", "month": "2026-13"}, headers=HEADERS
    )
    assert r.status_code == 422 and r.json()["error"]["details"]["fields"]["month"] == "invalid"
    r = client.get(
        "/v1/reports/pnl/export", params={"format": "xlsx", "month": "2027-01"}, headers=HEADERS
    )
    assert r.status_code == 422 and r.json()["error"]["details"]["fields"]["month"] == "in_future"


def test_too_many_exports_are_429_with_retry_after(client: TestClient) -> None:
    for _ in range(10):
        assert (
            client.get(
                "/v1/reports/stock/export", params={"format": "xlsx"}, headers=HEADERS
            ).status_code
            == 200
        )
    r = client.get("/v1/reports/stock/export", params={"format": "xlsx"}, headers=HEADERS)
    assert r.status_code == 429 and r.json()["error"]["code"] == "rate_limited"
    assert 1 <= int(r.headers["Retry-After"]) <= 61
    # an invalid request does not use the budget and is still validated first
    bad = client.get("/v1/reports/nothing/export", params={"format": "pdf"}, headers=HEADERS)
    assert bad.status_code == 404
    # JSON reports are not limited
    assert client.get("/v1/reports/stock", headers=HEADERS).status_code == 200


def test_published_openapi_has_the_report_endpoints_and_unit_pp() -> None:
    spec = json.loads(OPENAPI.read_text("utf-8"))
    assert "pp" in spec["components"]["schemas"]["FigureOut"]["properties"]["unit"]["enum"]
    for path in (
        "/v1/reports/pnl",
        "/v1/reports/norm",
        "/v1/reports/stock",
        "/v1/reports/stock/idle",
        "/v1/reports/{name}/export",
    ):
        assert path in spec["paths"]
